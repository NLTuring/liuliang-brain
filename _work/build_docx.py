# -*- coding: utf-8 -*-
"""构建《留量大脑》挑战杯参赛项目作品书（Word）。

运行：python build_docx.py
输出：留量大脑-电商赋能乡村的流量沉淀机制与数字治理平台-作品书.docx
"""
import os
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

import docx_common as D
from docx_common import (ACCENT, ACCENT2, DARK, GREY, TABLE_TOTAL_CM, CALLOUT_FILL,
                         ZEBRA_FILL)
import content_a as CA
import content_b as CB
import content_c as CC
import content_d as CD

PROJ = os.path.dirname(HERE)          # 项目根目录
FIGDIR = os.path.join(PROJ, "figures")
OUTNAME = "留量大脑-电商赋能乡村的流量沉淀机制与数字治理平台-作品书.docx"

HEADER_TEXT = "留量大脑 · 电商赋能乡村的流量沉淀机制与数字治理平台"


# ------------------------------------------------------------------ 封面
def build_cover(doc):
    cfg = CA.COVER
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    D.set_run(p.add_run(""), 8)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(6)
    D.set_run(p.add_run(cfg["kicker"]), 10, color=GREY)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    D.set_run(p.add_run(cfg["brand"]), 26, True, ACCENT)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    D.set_run(p.add_run(cfg["subtitle_en"]), 10, color=ACCENT2)

    # 分隔线
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.0
    D.set_run(p.add_run(""), 4)
    pPr = p._p.get_or_add_pPr()
    bd = OxmlElement("w:pBdr")
    bt = OxmlElement("w:bottom")
    bt.set(qn("w:val"), "single")
    bt.set(qn("w:sz"), "10")
    bt.set(qn("w:space"), "1")
    bt.set(qn("w:color"), "1B4D3E")
    bd.append(bt)
    pPr.append(bd)

    for i, line in enumerate(cfg["title_lines"]):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(1)
        D.set_run(p.add_run(line), 17, True, ACCENT)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(9)
    D.set_run(p.add_run(cfg["title_note"]), 10.5, color=ACCENT2)

    # 口号框
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    t.autofit = False
    cell = t.rows[0].cells[0]
    cell.text = ""
    D.cell_margins(t, top=95, bottom=95, left=150, right=150)
    D.cell_shade(cell, CALLOUT_FILL)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    D.set_run(p.add_run(cfg["slogan"]), 11.5, True, ACCENT)
    D.set_table_widths(t, [TABLE_TOTAL_CM])

    sp = doc.add_paragraph()
    sp.paragraph_format.space_after = Pt(0)
    sp.paragraph_format.line_spacing = 1.0
    D.set_run(sp.add_run(""), 6)

    # 信息表
    info = doc.add_table(rows=0, cols=2)
    info.style = "Table Grid"
    info.autofit = False
    D.cell_margins(info, top=42, bottom=42, left=120, right=120)
    for k, v in cfg["info"]:
        cells = info.add_row().cells
        D.no_split(info.rows[-1])
        for idx, txt in enumerate((k, v)):
            c = cells[idx]
            c.text = ""
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.22
            if idx == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                D.set_run(p.add_run(txt), 9.5, True, ACCENT)
                D.cell_shade(c, ZEBRA_FILL)
            else:
                D.set_run(p.add_run(txt), 9.5)
    D.set_table_widths(info, [3.3, 12.5])


# ------------------------------------------------------------------ 前置信息
def build_purpose(doc):
    for blk in CA.PURPOSE:
        render_block(doc, blk, pb_before=False)


def render_block(doc, blk, pb_before=False):
    kind = blk[0]
    if kind == "h1":
        D.H(doc, blk[1], 1, page_break_before=pb_before)
    elif kind == "h2":
        D.H(doc, blk[1], 2)
    elif kind == "h3":
        D.H(doc, blk[1], 3)
    elif kind == "p":
        D.para(doc, blk[1], indent=True)
    elif kind == "pn":
        p = D.para(doc, blk[1], indent=False)
        if len(blk) > 2 and blk[2]:
            D.keep_with_next(p)
    elif kind == "b":
        D.bullet(doc, blk[1])
    elif kind == "b2":
        D.bullet(doc, blk[1], level=1)
    elif kind == "num":
        D.numbered(doc, blk[1])
    elif kind == "quote":
        D.quote(doc, blk[1])
    elif kind == "table":
        D.table(doc, blk[1], blk[2], blk[3])
    elif kind == "caption":
        D.caption(doc, blk[1])
    elif kind == "tcap":
        # 表题：置于表格上方，并保证与表格同页
        D.caption(doc, blk[1], keep_next=True, as_title=True)
    elif kind == "callout":
        D.callout(doc, blk[1], blk[2])
    elif kind == "fig":
        D.figure(doc, blk[1], blk[2], blk[3] if len(blk) > 3 else 15.0, FIGDIR)
    elif kind == "pb":
        pass  # 分页统一由章节标题的 page_break_before 承担，避免产生空白页
    else:
        raise ValueError("unknown block: %r" % (kind,))


def reorder_captions(blocks):
    """把紧跟表格的表题提到表格之前（中文表格规范：表题在上）。"""
    out = []
    i = 0
    while i < len(blocks):
        b = blocks[i]
        if (b[0] == "table" and i + 1 < len(blocks) and blocks[i + 1][0] == "caption"):
            out.append(("tcap", blocks[i + 1][1]))
            out.append(b)
            i += 2
            continue
        out.append(b)
        i += 1
    return out


def render_blocks(doc, blocks, pb_before_h1=False):
    for blk in reorder_captions(blocks):
        pb = pb_before_h1 and blk[0] == "h1" and blk[1] in PAGE_BREAK_HEADS
        render_block(doc, blk, pb_before=pb)


# ------------------------------------------------------------------ 目录页
ALL_BLOCKS = [CA.BLOCKS_A, CB.BLOCKS_B, CC.BLOCKS_C, CD.BLOCKS_D]

# 只在正文起始处分页；其余章节与附录连续排版，避免章节末出现半空页
PAGE_BREAK_HEADS = {"第一章  项目速览"}


def collect_headings():
    """按正文顺序收集标题，供生成目录。"""
    hs = []
    for blocks in ALL_BLOCKS:
        for b in blocks:
            if b[0] in ("h1", "h2", "h3"):
                hs.append((int(b[0][1]), b[1]))
    return hs


def load_page_map():
    import json
    path = os.path.join(HERE, "page_map.json")
    if os.path.exists(path):
        with open(path, encoding="utf-8") as f:
            return json.load(f)
    print("  (page_map.json 不存在：目录页码先用占位，渲染后再回填)")
    return {}


def build_toc_page(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.page_break_before = True
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    D.set_run(p.add_run("目　录"), 18, True, ACCENT)

    hint = doc.add_paragraph()
    hint.alignment = WD_ALIGN_PARAGRAPH.CENTER
    hint.paragraph_format.space_after = Pt(10)
    D.set_run(hint.add_run(CA.TOC_HINT), 8.5, color=GREY)

    pmap = load_page_map()
    entries = [(lvl, text, pmap.get(text, "—")) for lvl, text in collect_headings()]
    D.add_toc_field(doc, entries, "1-3")


# ------------------------------------------------------------------ 主流程
def main():
    doc = D.new_document()

    # --- 第 1 节：封面（无页眉页脚）---
    build_cover(doc)

    # --- 第 2 节：正文（含前置信息、目录、正文、附录）---
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    D.setup_page(sec)
    sec.header.is_linked_to_previous = False
    sec.footer.is_linked_to_previous = False
    D.add_body_header(sec, HEADER_TEXT)
    D.add_body_footer(sec)

    build_purpose(doc)
    build_toc_page(doc)

    render_blocks(doc, CA.BLOCKS_A, pb_before_h1=True)
    render_blocks(doc, CB.BLOCKS_B, pb_before_h1=True)
    render_blocks(doc, CC.BLOCKS_C, pb_before_h1=True)
    render_blocks(doc, CD.BLOCKS_D, pb_before_h1=True)

    out = os.path.join(PROJ, OUTNAME)
    doc.save(out)
    print("saved:", out)
    return out


if __name__ == "__main__":
    main()
